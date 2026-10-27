"""Duitse bronnen voor de TAM/whitespace-lijstbouwer: Destatis GV, Destatis FEU en handtabellen.

Destatis Gemeindeverzeichnis, kwartaaluittreksel (GVAuszugQ)
------------------------------------------------------------
URL zonder jaartal: .../GVAuszugQ/AuszugGV<Q>QAktuell.xlsx (Q = 1..4); alle vier worden
opgehaald en het bestand met de nieuwste Gebietsstand wint. Vrij herbruikbaar met
bronvermelding ("Vervielfaeltigung und Verbreitung ... mit Quellenangabe gestattet").
Geinspecteerd 2026-10-27: Q1 31.03.2026, Q2 30.06.2026, Q3 30.09.2026 (nieuwste), Q4 31.12.2025.
Werkblad 2 'Onlineprodukt_Gemeinden<ddmmjjjj>'; rij 1 titel '... am 30.09.2026 (3. Quartal)',
rij 3 kop (Satzart, Textkennzeichen, Amtlicher Regionalschluessel (ARS), Gemeindename,
Flaeche, 'Bevoelkerung auf Grundlage des Zensus 2022', Postleitzahl, ...), rij 4 subkop
(Land, RB, Kreis, VB, Gem; insgesamt/maennlich/weiblich), rij 5 peildata (Gebietsstand;
bevolking '31.12.2025 (Jahr)'), data vanaf rij 7, voetnoten onderaan. De kop wordt op tekst
gezocht, niet op rijnummer. Sleutelvelden zijn tekst: Land 2, RB 1, Kreis 2, VB 4, Gem 3 cijfers.
  AGS (8) = Land+RB+Kreis+Gem; Verbandsschluessel (9) = Land+RB+Kreis+VB; ARS (12) = alle vijf.

Verdeling Satzart / Textkennzeichen (Q3 2026, 15.984 datarijen):
  10 Land                                  16
  20 RB                                    29  (echt: 19 in Land 05 (5), 06 (3), 08 (4), 09 (7);
                                                niet: 03 'Statistische Region' (4), 07 en 14 'frueher:')
  30 Region (alleen BW)                    12  genegeerd
  40 Kreis   41 kreisfreie Stadt 98, 42 Stadtkreis 9, 43 Kreis 42, 44 Landkreis 251,
             45 Regionalverband 1          => 401 Kreise, 107 kreisfrei, 294 Landkreise
  50 Verband 50 verbandsfreie Gemeinde (placeholder) 3.381, 51 Amt 207, 52 Samtgemeinde 114,
             53 Verbandsgemeinde 148, 54 Verwaltungsgemeinschaft 573, 55 Kirchspielslandgemeinde 2,
             56 Verwaltungsverband 120, 58 Erfuellende Gemeinde 38  => 1.202 echte Verbaende
  60 Gemeinde 60 Markt 386, 61 kreisfreie Stadt 98, 62 Stadtkreis 9, 63 Stadt 1.828,
             64 kreisangehoerige Gemeinde 8.301, 67 Grosse Kreisstadt 125 => 10.747 Gemeinden;
             65 bewohntes (2) en 66 unbewohntes (194) gemeindefreies Gebiet => geen entiteit.
  Correctie op de brief: gemeindefreie Gebiete zijn 65/66, niet 63 (63 = Stadt).
  Bevolking staat alleen op Gemeinde-rijen; som 10.747 Gemeinden = 83.466.003 (31.12.2025).
  Gem '000' komt precies voor bij de 107 kreisfreie Staedte (61/62).
Betekenis van de Textkennzeichen: Destatis, 'Beschreibung der in GV-ISys enthaltenen Regional-
und Gebietseinheiten' (2020-09-02), p. 22-25:
https://www.destatis.de/DE/Themen/Laender-Regionen/Regionales/Gemeindeverzeichnis/Administrativ/beschreibung-gebietseinheiten.pdf

Verband-typen per Land (land, tk: aantal): 01 SH 51:81, 55:2 | 03 NI 52:114 | 07 RP 53:129 |
08 BW 54:156 (VVG), 56:114 (GVV) | 09 BY 54:311 | 12 BB 51:50, 53:1, 58:1 | 13 MV 51:76 |
14 SN 54:64, 56:6 | 15 ST 53:18 | 16 TH 54:42, 58:37.
Tk 58 (TH 37, BB 1 = mitverwaltende Gemeinde Schwedt/Oder): de naam van de Verband-rij is in
alle 38 gevallen exact gelijk aan de naam van precies een lid; zo wordt de erfuellende Gemeinde
uit de GV-data herkend. SN tk 54 (Verwaltungsgemeinschaft, erfuellende-Gemeinde-model volgens
SaechsKomZG § 36) markeert de erfuellende Gemeinde NIET (8/64 namen exact gelijk); daar wordt
op kernnaam gezocht (36/64 eenduidig, Q3 2026); de rest blijft Verband-entiteit met WARNING.

Destatis FEU-lijsten (Fonds, Einrichtungen, Unternehmen)
--------------------------------------------------------
.../Fonds-Einrichtungen-Unternehmen/Methoden/Downloads/liste-extrahaushalte-<jaar>-xlsx.xlsx en
liste-sonstige-FEU-<jaar>-xlsx.xlsx (2026: 8.073 + 17.681 rijen). Blad 3 (Liste_der_...),
kop 'Identifikator, Berichtsstelle, [Teilsektor,] Ebene, Land' op rij 2. Identifikator is
15 cijfers (tekst, kan met 0 beginnen). 2.452 rijen met 'zweckverband' in de naam, geen
dubbele ids tussen de lijsten, allemaal Ebene 'Gemeinde/Gemeindeverbaende'. Stabiliteit:
2.438 van die ids staan ook in de 2025-lijsten, 99% met dezelfde naam -> id DE-FEU-<id>.
"""
from __future__ import annotations

import logging
import re
from datetime import date

import openpyxl
import pandas as pd

from step2_enrich import normalize
from tam.entities import SourceBundle, SourceStatus, new_entity

log = logging.getLogger(__name__)

GV_URL = ('https://www.destatis.de/DE/Themen/Laender-Regionen/Regionales/Gemeindeverzeichnis/'
          'Administrativ/Archiv/GVAuszugQ/AuszugGV{q}QAktuell.xlsx?__blob=publicationFile')
GV_SOURCE = 'Destatis Gemeindeverzeichnis'
FEU_URL = ('https://www.destatis.de/DE/Themen/Staat/Oeffentliche-Finanzen/'
           'Fonds-Einrichtungen-Unternehmen/Methoden/Downloads/{file}?__blob=publicationFile')
FEU_FILES = {'extra': 'liste-extrahaushalte-{year}-xlsx.xlsx',
             'sonst': 'liste-sonstige-FEU-{year}-xlsx.xlsx'}
FEU_SOURCE = 'Destatis FEU-Liste'

GV_COLUMNS = ['satzart', 'textkennzeichen', 'land', 'rb', 'kreis', 'vb', 'gem', 'name',
              'population', 'plz']

# Stadstaten: Land -> AGS van de Gemeinde die de Land-entiteit vertegenwoordigt.
# Bremen: alleen Stadt Bremen; Bremerhaven heeft een eigen bestuur (kreisfreie Stadt).
CITY_STATES = {'02': '02000000', '11': '11000000', '04': '04011000'}
# Laender met echte Regierungsbezirke (Satzart 20); 03 = Statistische Regionen, 07/14 'frueher'
RB_LAENDER = {'05', '06', '08', '09'}
KREISFREI_TK = {'41', '42'}
GEMEINDEFREI_TK = {'65', '66'}

# Gemeindeverbaende per (Land, Textkennzeichen) -> (regel, label).
#   verband     Verband is de entiteit, leden zijn aliassen, bevolking = som
#   erfuellend  de erfuellende Gemeinde is de entiteit met bevolking van de hele Verband,
#               overige leden zijn aliassen, geen Verband-entiteit (herkend op exacte naam)
#   erfuellend_core  idem, herkend op kernnaam; zonder eenduidige treffer -> 'verband'
#   gvv         Verband is entiteit (eigen personeel), leden blijven ook entiteit
#   vvg         geen Verband-entiteit, leden blijven entiteit
# Onbekende combinaties: Verband-entiteit + leden blijven entiteit + WARNING.
VERBAND_RULES = {
    # SH Amtsordnung (Aemter, ook de twee Kirchspielslandgemeinden in Dithmarschen zijn Aemter):
    # https://www.gesetze-rechtsprechung.sh.juris.de/bssh/document/jlr-AmtsOSH2003rahmen
    ('01', '51'): ('verband', 'Amt'),
    ('01', '55'): ('verband', 'Kirchspielslandgemeinde'),
    # NI NKomVG §§ 97 ff. Samtgemeinden: https://voris.wolterskluwer-online.de/browse/document/2c9b8b1c-60c4-3a62-8c1d-0f1a2e8a0d0b
    ('03', '52'): ('verband', 'Samtgemeinde'),
    # RP GemO §§ 64 ff. Verbandsgemeinden: https://landesrecht.rlp.de/bsrp/document/jlr-GemORPrahmen
    ('07', '53'): ('verband', 'Verbandsgemeinde'),
    # BW GemO §§ 59-62: Gemeindeverwaltungsverband (eigen Koerperschaft) vs. vereinbarte
    # Verwaltungsgemeinschaft (leden houden eigen Rathaus):
    # https://www.landesrecht-bw.de/bsbw/document/jlr-GemOBW2000pP59
    ('08', '56'): ('gvv', 'Gemeindeverwaltungsverband'),
    ('08', '54'): ('vvg', 'Vereinbarte Verwaltungsgemeinschaft'),
    # BY Verwaltungsgemeinschaftsordnung (VGemO): https://www.gesetze-bayern.de/Content/Document/BayVwGemO
    ('09', '54'): ('verband', 'Verwaltungsgemeinschaft'),
    # BB Amtsordnung (BbgKVerf §§ 133 ff.) en Verbandsgemeinde- und Mitverwaltungsgesetz (VgMvG):
    # https://bravors.brandenburg.de/gesetze/bbgkverf  https://bravors.brandenburg.de/gesetze/vgmvg
    ('12', '51'): ('verband', 'Amt'),
    ('12', '53'): ('verband', 'Verbandsgemeinde'),
    ('12', '58'): ('erfuellend', 'mitverwaltende Gemeinde'),
    # MV Kommunalverfassung §§ 125 ff. Aemter: https://www.landesrecht-mv.de/bsmv/document/jlr-KVMV2011rahmen
    ('13', '51'): ('verband', 'Amt'),
    # SN SaechsKomZG: Verwaltungsverband (§§ 3 ff.) is eigen Koerperschaft; Verwaltungsgemeinschaft
    # (§ 36) = erfuellende Gemeinde. De GV markeert die niet: herkend op kernnaam (naam tot de
    # eerste komma, genormaliseerd) gelijk aan precies een lid; anders Verband-entiteit + WARNING.
    # https://www.revosax.sachsen.de/vorschrift/2649-SaechsKomZG
    ('14', '56'): ('verband', 'Verwaltungsverband'),
    ('14', '54'): ('erfuellend_core', 'Verwaltungsgemeinschaft'),
    # ST KVG LSA §§ 88 ff. Verbandsgemeinden: https://www.landesrecht.sachsen-anhalt.de/bsst/document/jlr-KomVerfGSTrahmen
    ('15', '53'): ('verband', 'Verbandsgemeinde'),
    # TH ThuerKO § 46 Verwaltungsgemeinschaft, § 51 erfuellende Gemeinde:
    # https://landesrecht.thueringen.de/bsth/document/jlr-KomOTHrahmen
    ('16', '54'): ('verband', 'Verwaltungsgemeinschaft'),
    ('16', '58'): ('erfuellend', 'erfuellende Gemeinde'),
}
# Afkortingen in GV-namen van Verbaende die door het label vervangen worden
_VERBAND_NAME_STRIP = [re.compile(r'\s*\(VGem\)\s*$'), re.compile(r'^GVV\s+'), re.compile(r'^VVG\s+')]


def _s(value) -> str:
    """Celwaarde als getrimde tekst; leeg bij None/NaN."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    return re.sub(r'\s+', ' ', str(value)).strip()


def match_name_de(name: str) -> str:
    """Officiele naam tot de eerste komma: 'Kassel, Stadt' -> 'Kassel'."""
    return name.split(',', 1)[0].strip()


def _verband_core(name: str) -> str:
    core = name
    for rx in _VERBAND_NAME_STRIP:
        core = rx.sub('', core)
    return core.strip()


def _kreis_display(name: str, tk: str) -> str:
    """Landkreis/Kreis-prefix als de GV-naam het type niet al draagt."""
    low = name.lower()
    if tk == '45' or 'kreis' in low or low.startswith(('region ', 'städteregion')):
        return name
    return f"{'Kreis' if tk == '43' else 'Landkreis'} {name}"


def _rb_core(name: str) -> str:
    return re.sub(r'^Reg\.-Bez\.\s*', '', name).strip()


def parse_gv(df: pd.DataFrame, stichtag: str, population_date: str = '') -> SourceBundle:
    """Bouw DE-entiteiten uit een genormaliseerd GV-blad (kolommen GV_COLUMNS)."""
    d = df.copy()
    for c in GV_COLUMNS:
        if c != 'population':
            d[c] = d[c].map(_s)
    d = d[d['satzart'].isin(['10', '20', '40', '50', '60'])]
    d['k5'] = d['land'] + d['rb'] + d['kreis']
    d['vkey'] = d['k5'] + d['vb']
    d['ags'] = d['k5'] + d['gem']

    gem = d[(d['satzart'] == '60') & ~d['textkennzeichen'].isin(GEMEINDEFREI_TK)].copy()
    gem['pop'] = pd.to_numeric(gem['population'], errors='coerce')
    missing = gem[gem['pop'].isna()]
    if len(missing):
        log.warning('GV: %d Gemeinden zonder bevolking (blijft leeg): %s', len(missing),
                    ', '.join(missing['name']))

    def pop_of(value):
        return None if pd.isna(value) else int(value)

    def pop_sum(sub):
        known = sub['pop'].dropna()
        return int(known.sum()) if len(known) else None

    def sum_fields(sub):
        total = pop_sum(sub)
        return {'population': total,
                'population_method': 'sum of members' if total is not None else 'unknown'}

    land_name = {r.land: r.name for r in d[d['satzart'] == '10'].itertuples()}
    kreis_rows = d[d['satzart'] == '40']
    kreis_name = {r.k5: r.name for r in kreis_rows.itertuples()}
    kreisfrei = {r.k5 for r in kreis_rows.itertuples() if r.textkennzeichen in KREISFREI_TK}
    gem_by_vkey = {k: g for k, g in gem.groupby('vkey')}
    gem_ags_set = set(gem['ags'])
    city_gem = set(CITY_STATES.values())

    def prefix_members(prefix):
        return gem[gem['ags'].str.startswith(prefix)]

    common = dict(country='DE', source=GV_SOURCE, source_date=stichtag,
                  population_source=f'{GV_SOURCE} {stichtag}', population_date=population_date)
    rows, members, aliases = [], [], []
    gem_override: dict[str, dict] = {}   # ags -> velden voor erfuellende Gemeinde
    aliased: set[str] = set()
    counts: dict[str, int] = {}

    def gem_kind(ags):
        return 'de_kreisfreie_stadt' if ags[:5] in kreisfrei else 'de_gemeinde'

    def territorial(eid, ags_list):
        for a in ags_list:
            if a in city_gem:
                continue
            members.append({'entity_id': eid, 'member_entity_id': f'DE-AGS-{a}',
                            'member_kind': gem_kind(a), 'membership_type': 'territorial',
                            'join_date': ''})

    # Laender en stadstaten
    for r in d[d['satzart'] == '10'].itertuples():
        eid = f'DE-AGS-{r.land}'
        if r.land in CITY_STATES:
            own = gem[gem['ags'] == CITY_STATES[r.land]]
            if len(own) != 1:
                log.warning('GV: stadstaat %s zonder Gemeinde %s', r.name, CITY_STATES[r.land])
                continue
            g = own.iloc[0]
            rows.append(new_entity(entity_id=eid, name=r.name, kind='de_stadtstaat',
                                   population=pop_of(g['pop']),
                                   population_method='own' if pop_of(g['pop']) is not None else 'unknown',
                                   postcode=g['plz'], official_ids=f"ars:{g['k5']}{g['vb']}{g['gem']}",
                                   **common))
            continue
        sub = prefix_members(r.land)
        rows.append(new_entity(entity_id=eid, name=r.name, kind='de_bundesland', **sum_fields(sub),
                               official_ids=f'ars:{r.land}', **common))
        territorial(eid, sub['ags'])

    # Regierungsbezirke
    for r in d[d['satzart'] == '20'].itertuples():
        if r.land not in RB_LAENDER:
            log.info('GV: RB-rij overgeslagen (geen bestuurslaag): %s', r.name)
            continue
        key = r.land + r.rb
        sub = prefix_members(key)
        eid = f'DE-AGS-{key}'
        core = _rb_core(r.name)
        rows.append(new_entity(entity_id=eid, name=f'Regierungsbezirk {core}', match_name=core,
                               kind='de_regierungsbezirk', region=land_name.get(r.land, ''),
                               **sum_fields(sub), official_ids=f'ars:{key}', **common))
        territorial(eid, sub['ags'])

    # Landkreise (kreisfreie Staedte worden via hun Gemeinde-rij een entiteit)
    for r in kreis_rows.itertuples():
        if r.textkennzeichen in KREISFREI_TK:
            continue
        if r.textkennzeichen not in {'43', '44', '45'}:
            log.warning('GV: onbekend Kreis-Textkennzeichen %s voor %s; als Landkreis behandeld',
                        r.textkennzeichen, r.name)
        sub = prefix_members(r.k5)
        eid = f'DE-AGS-{r.k5}'
        rows.append(new_entity(entity_id=eid, name=_kreis_display(r.name, r.textkennzeichen),
                               match_name=match_name_de(r.name), kind='de_landkreis',
                               region=land_name.get(r.land, ''), **sum_fields(sub),
                               official_ids=f'ars:{r.k5}', **common))
        territorial(eid, sub['ags'])

    # Gemeindeverbaende
    unknown_types: dict[tuple, list] = {}
    verband_members: dict[str, list] = {}
    alias_ids: dict[str, str] = {}
    core_unresolved: list[str] = []
    core_resolved = 0

    def add_alias(member, target_id, reason):
        aliased.add(member.ags)
        alias_ids[f'DE-AGS-{member.ags}'] = target_id
        aliases.append({'alias_name': match_name_de(member.name), 'country': 'DE',
                        'entity_id': target_id, 'reason': reason})
    for r in d[d['satzart'] == '50'].itertuples():
        if r.textkennzeichen == '50':
            continue  # 'verbandsfreie Gemeinde': placeholder
        sub = gem_by_vkey.get(r.vkey)
        if sub is None or not len(sub):
            log.warning('GV: Verband %s (%s) zonder Gemeinden', r.name, r.vkey)
            continue
        verband_members[f'DE-ARS-{r.vkey}'] = [f'DE-AGS-{a}' for a in sub['ags']]
        rule, label = VERBAND_RULES.get((r.land, r.textkennzeichen), (None, None))
        if rule is None:
            unknown_types.setdefault((r.land, r.textkennzeichen), []).append(r.name)
            rule, label = 'gvv', 'Gemeindeverband'
        if rule in ('erfuellend', 'erfuellend_core'):
            if rule == 'erfuellend':
                hit = sub[sub['name'] == r.name]
            else:
                core = normalize(match_name_de(r.name))
                hit = sub[sub['name'].map(lambda n: normalize(match_name_de(n))) == core]
            if len(hit) != 1:
                if rule == 'erfuellend':
                    log.warning('GV: erfuellende Gemeinde voor %s (%s) niet eenduidig; als Verband '
                                'behandeld', r.name, r.vkey)
                else:
                    core_unresolved.append(f'{r.name} ({r.vkey})')
                rule = 'verband'
            else:
                core_resolved += rule == 'erfuellend_core'
                eg = hit.iloc[0]
                gem_override[eg['ags']] = sum_fields(sub)
                for m in sub[sub['ags'] != eg['ags']].itertuples():
                    add_alias(m, f"DE-AGS-{eg['ags']}", f"administered by {eg['name']}")
                counts[f'{label} (erfuellend)'] = counts.get(f'{label} (erfuellend)', 0) + 1
                continue
        if rule == 'vvg':
            counts[label] = counts.get(label, 0) + 1
            continue
        core = _verband_core(r.name)
        if core.lower().startswith(label.lower() + ' '):
            name, core = core, core[len(label):].strip()
        else:
            name = f'{label} {core}'
        eid = f'DE-ARS-{r.vkey}'
        rows.append(new_entity(entity_id=eid, name=name, match_name=match_name_de(core),
                               kind='de_gemeindeverband', region=kreis_name.get(r.k5, ''),
                               **sum_fields(sub), official_ids=f'ars:{r.vkey}', **common))
        counts[label] = counts.get(label, 0) + 1
        for m in sub.itertuples():
            if rule == 'verband':
                add_alias(m, eid, f'member of {name}')
            else:
                members.append({'entity_id': eid, 'member_entity_id': f'DE-AGS-{m.ags}',
                                'member_kind': 'de_gemeinde', 'membership_type': 'cooperation',
                                'join_date': ''})
    if core_unresolved:
        log.warning('GV: %d Verwaltungsgemeinschaften (SN) zonder eenduidige erfuellende Gemeinde op '
                    'kernnaam; als Verband-entiteit behandeld: %s', len(core_unresolved),
                    '; '.join(core_unresolved))
    for (land, tk), names in unknown_types.items():
        log.warning('GV: Verband-type %s in Land %s niet in VERBAND_RULES (%d): %s; Verband en leden '
                    'blijven entiteit', tk, land, len(names), ', '.join(names[:10]))

    # Gemeinden (incl. kreisfreie Staedte), behalve aliassen en stadstaat-Gemeinden
    for g in gem.itertuples():
        if g.ags in aliased or g.ags in city_gem:
            continue
        kind = gem_kind(g.ags)
        region = land_name.get(g.land, '') if kind == 'de_kreisfreie_stadt' else kreis_name.get(g.k5, '')
        own_pop = pop_of(g.pop)
        fields = {'population': own_pop,
                  'population_method': 'own' if own_pop is not None else 'unknown'}
        fields.update(gem_override.get(g.ags, {}))
        rows.append(new_entity(entity_id=f'DE-AGS-{g.ags}', name=g.name,
                               match_name=match_name_de(g.name), kind=kind, region=region,
                               postcode=g.plz, official_ids=f'ars:{g.vkey}{g.gem}',
                               **fields, **common))

    # Gemeinden zonder bevolking blijven buiten municipal_population
    muni_pop = {f'DE-AGS-{g.ags}': int(g.pop) for g in gem.itertuples()
                if g.ags not in city_gem and not pd.isna(g.pop)}
    for land, ags in CITY_STATES.items():
        if ags in gem_ags_set:
            value = pop_of(gem.loc[gem['ags'] == ags, 'pop'].iloc[0])
            if value is not None:
                muni_pop[f'DE-AGS-{land}'] = value

    bundle = SourceBundle.empty()
    bundle.entities = pd.DataFrame(rows, columns=bundle.entities.columns)
    bundle.memberships = pd.DataFrame(members, columns=bundle.memberships.columns)
    bundle.aliases = pd.DataFrame(aliases, columns=bundle.aliases.columns)
    bundle.extras['municipal_population'] = muni_pop
    # Gemeindeverband -> lid-Gemeinden, voor handtabellen met een Amt als lid (load_de_hand)
    bundle.extras['de_verband_members'] = verband_members
    # Alias-Gemeinde (DE-AGS-<8>) -> entiteit die haar vertegenwoordigt (R12)
    bundle.extras['alias_ids'] = alias_ids
    bundle.extras['de_population_date'] = population_date or stichtag
    kinds = bundle.entities['kind'].value_counts().to_dict()
    msg = (f"{len(gem)} Gemeinden ({len(aliased)} als alias), "
           + ', '.join(f'{k} {v}' for k, v in sorted(kinds.items()))
           + '; Verbaende: ' + ', '.join(f'{k} {v}' for k, v in sorted(counts.items()))
           + f'; SN-VG erfuellend op kernnaam {core_resolved}, als Verband {len(core_unresolved)}')
    log.info('GV: %s', msg)
    bundle.status.append(SourceStatus('de_gv', True, len(bundle.entities), stichtag, msg))
    return bundle


# ---------------------------------------------------------------------------
# Inlezen van het GV-xlsx
# ---------------------------------------------------------------------------
_DATE_RE = re.compile(r'(\d{2})\.(\d{2})\.(\d{4})')


def _iso(m) -> str:
    return f'{m.group(3)}-{m.group(2)}-{m.group(1)}'


def _gv_sheet(path):
    """Naam van het datablad: het blad met 'Satzart' in de eerste rijen."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=10, values_only=True):
                if any(_s(v) == 'Satzart' for v in row):
                    return ws.title
    finally:
        wb.close()
    raise ValueError(f'Geen GV-datablad gevonden in {path}')


def gv_stichtag(path) -> str:
    """Gebietsstand (ISO) uit de titelrij van het datablad."""
    sheet = _gv_sheet(path)
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        for row in wb[sheet].iter_rows(max_row=6, values_only=True):
            for v in row:
                m = _DATE_RE.search(_s(v))
                if m:
                    return _iso(m)
    finally:
        wb.close()
    raise ValueError(f'Geen Gebietsstand gevonden in {path}')


def read_gv_xlsx(path) -> tuple[pd.DataFrame, str]:
    """(genormaliseerd frame met GV_COLUMNS, bevolkingspeildatum ISO) uit een GV-xlsx."""
    raw = pd.read_excel(path, sheet_name=_gv_sheet(path), header=None, dtype=str)
    cells = raw.map(_s)
    hdr = next(i for i in range(min(20, len(cells))) if 'Satzart' in set(cells.iloc[i]))
    top, sub = cells.iloc[hdr], cells.iloc[hdr + 1]

    def col(row, pred):
        hits = [j for j, v in enumerate(row) if pred(v)]
        if not hits:
            raise ValueError('GV-kolom niet gevonden')
        return hits[0]

    pos = {
        'satzart': col(top, lambda v: v == 'Satzart'),
        'textkennzeichen': col(top, lambda v: v == 'Textkennzeichen'),
        'land': col(sub, lambda v: v == 'Land'), 'rb': col(sub, lambda v: v == 'RB'),
        'kreis': col(sub, lambda v: v == 'Kreis'), 'vb': col(sub, lambda v: v == 'VB'),
        'gem': col(sub, lambda v: v == 'Gem'),
        'name': col(top, lambda v: v.startswith('Gemeindename')),
        'plz': col(top, lambda v: v.startswith('Postleitzahl')),
    }
    pop_top = col(top, lambda v: v.startswith('Bevölkerung'))
    pos['population'] = next(j for j in range(pop_top, len(sub)) if sub.iloc[j] == 'insgesamt')
    pop_date = ''
    for i in range(hdr + 2, hdr + 4):
        m = _DATE_RE.search(cells.iloc[i, pos['population']])
        if m:
            pop_date = _iso(m)
            break
    data = raw.iloc[hdr + 2:, [pos[c] for c in GV_COLUMNS]].copy()
    data.columns = GV_COLUMNS
    data = data[data['satzart'].map(_s).str.fullmatch(r'\d{2}')]
    return data.reset_index(drop=True), pop_date


def gv_candidates(cache) -> list[tuple[str, int, object]]:
    """(Gebietsstand, kwartaal, pad) van alle beschikbare kwartaalbestanden, nieuwste eerst."""
    found = []
    for q in range(1, 5):
        try:
            path = cache.fetch(f'de_gv_q{q}', GV_URL.format(q=q), suffix='.xlsx', max_age_days=90)
            found.append((gv_stichtag(path), q, path))
        except Exception as e:  # noqa: BLE001
            log.warning('GV kwartaal %d niet beschikbaar: %s', q, e)
    return sorted(found, key=lambda t: (t[0], t[1]), reverse=True)


def load_de_gv(cache) -> SourceBundle:
    """Verwerk het kwartaalbestand met de nieuwste Gebietsstand; faalt dat, dan het volgende."""
    candidates = gv_candidates(cache)
    errors = []
    for stichtag, q, path in candidates:
        try:
            df, pop_date = read_gv_xlsx(path)
            bundle = parse_gv(df, stichtag, population_date=pop_date)
        except Exception as e:  # noqa: BLE001
            log.error('GV kwartaal %d (%s) verwerken mislukt: %s; probeer volgende', q, stichtag, e)
            errors.append(f'Q{q} {stichtag}: {e}')
            continue
        log.info('GV: Gebietsstand %s (kwartaal %d) gebruikt', stichtag, q)
        return bundle
    bundle = SourceBundle.empty()
    msg = 'mislukt: ' + ('; '.join(errors) if errors else 'geen GV-bestand')
    bundle.status.append(SourceStatus('de_gv', False, 0, '', msg))
    return bundle


# ---------------------------------------------------------------------------
# FEU-lijsten
# ---------------------------------------------------------------------------
def parse_feu(frames: list[pd.DataFrame], it_names, source_date: str) -> SourceBundle:
    """Zweckverbaende uit FEU-frames (kolommen id, name, ebene, land); IT-handtabel eruit."""
    it_norm = {normalize(n) for n in it_names}
    seen, rows, n_it = set(), [], 0
    for f in frames:
        for r in f.itertuples():
            fid, name = _s(r.id), _s(r.name)
            if 'zweckverband' not in name.lower() or not fid or fid in seen:
                continue
            seen.add(fid)
            if normalize(name) in it_norm:
                n_it += 1
                continue
            rows.append(new_entity(
                entity_id=f'DE-FEU-{fid}', name=name, match_name=name, country='DE',
                kind='de_zweckverband', region=_s(r.land), population_method='unknown',
                source=f'{FEU_SOURCE} {source_date[:4]}', source_date=source_date))
    bundle = SourceBundle.empty()
    bundle.entities = pd.DataFrame(rows, columns=bundle.entities.columns)
    msg = f'{len(rows)} Zweckverbaende ({n_it} via IT-handtabel overgeslagen)'
    log.info('FEU: %s', msg)
    bundle.status.append(SourceStatus('de_feu', True, len(rows), source_date, msg))
    return bundle


def read_feu_xlsx(path) -> pd.DataFrame:
    """Lees het lijstblad van een FEU-xlsx (kop 'Identifikator' op tekst gezocht)."""
    with pd.ExcelFile(path) as xl:  # sluit het werkboek (geen ResourceWarning)
        for sheet in xl.sheet_names:
            if not sheet.strip().startswith('Liste'):
                continue
            raw = xl.parse(sheet, header=None, dtype=str)
            first = raw.iloc[:, 0].map(_s)
            hdr = first[first == 'Identifikator'].index[0]
            header = [_s(v) for v in raw.loc[hdr]]
            df = raw.loc[hdr + 1:].copy()
            df.columns = header
            df = df[df['Identifikator'].map(_s) != '']
            return pd.DataFrame({'id': df['Identifikator'].map(_s), 'name': df['Berichtsstelle'].map(_s),
                                 'ebene': df['Ebene'].map(_s), 'land': df['Land'].map(_s)})
    raise ValueError(f'Geen FEU-lijstblad in {path}')


def load_de_feu(cache, today: date | None = None) -> SourceBundle:
    """Beide FEU-lijsten van het nieuwste beschikbare jaar (dit jaar, anders vorig jaar)."""
    from tam.hand_tables_de import it_feu_names
    today = today or date.today()
    last_error = None
    for year in (today.year, today.year - 1):
        try:
            frames = []
            for key, pattern in FEU_FILES.items():
                path = cache.fetch(f'de_feu_{key}_{year}', FEU_URL.format(file=pattern.format(year=year)),
                                   suffix='.xlsx')
                frames.append(read_feu_xlsx(path))
            return parse_feu(frames, it_feu_names(), f'{year}-01-01')
        except Exception as e:  # noqa: BLE001
            last_error = e
            log.warning('FEU-lijsten %d niet beschikbaar: %s', year, e)
    bundle = SourceBundle.empty()
    bundle.status.append(SourceStatus('de_feu', False, 0, '', f'mislukt: {last_error}'))
    return bundle


# ---------------------------------------------------------------------------
# Handtabellen
# ---------------------------------------------------------------------------
def _key_population(keys, municipal_population, verband_members=None):
    """Som over alle verschillende DE-gemeenten onder de keys (AGS-prefix of Verband-id).

    Geeft (som, aantal gemeenten, onopgeloste keys)."""
    verband_members = verband_members or {}
    found, unresolved = set(), []
    for key in keys:
        if key.startswith('DE-ARS-'):
            ids = [m for m in verband_members.get(key, []) if m in municipal_population]
        else:
            prefix = key.removeprefix('DE-AGS-')
            ids = [m for m in municipal_population
                   if m.startswith('DE-AGS-') and m[7:].startswith(prefix)]
        if not ids:
            unresolved.append(key)
        found.update(ids)
    return sum(municipal_population[m] for m in found), len(found), unresolved


def _municipal_keys(keys, municipal_population):
    """Alleen gemeentelijke leden (spec 3.4): Gemeinden, kreisfreie Staedte, stadstaten en
    Gemeindeverbaende (= hun lid-Gemeinden). Een Kreis/RB/Land als lid telt niet mee.
    Een 5-cijferige sleutel van een kreisfreie Stadt telt als die Stadt (<k5>000)."""
    kept, skipped = [], []
    for key in keys:
        bare = key.removeprefix('DE-AGS-')
        if key.startswith('DE-ARS-') or len(bare) == 8 or bare in CITY_STATES:
            kept.append(key)
        elif len(bare) == 5 and f'DE-AGS-{bare}000' in municipal_population:
            kept.append(f'{bare}000')
        else:
            skipped.append(key)
    return kept, skipped


def _member_id(key: str) -> str:
    return key if key.startswith('DE-') else f'DE-AGS-{key}'


def _member_kind(key: str) -> str:
    if key.startswith('DE-ARS-'):
        return 'de_gemeindeverband'
    key = key.removeprefix('DE-AGS-')
    if len(key) == 2:
        return 'de_stadtstaat' if key in CITY_STATES else 'de_bundesland'
    if len(key) == 3:
        return 'de_regierungsbezirk'
    if len(key) == 5:
        return 'de_landkreis'
    return 'de_kreisfreie_stadt' if key.endswith('000') else 'de_gemeinde'


def load_de_hand(municipal_population: dict, verband_members: dict | None = None,
                 population_date: str = '') -> SourceBundle:
    """DE-handtabellen: Bezirke, hoehere Kommunalverbaende, IT-Zweckverbaende, Bundesministerien.

    `verband_members` = extras['de_verband_members'] uit parse_gv; nodig voor leden die een
    Gemeindeverband zijn (bijv. een Amt in de kommunit-satzung). `population_date` =
    extras['de_population_date'] (peildatum van de GV-bevolking waarover gesommeerd wordt).
    Hoehere Kommunalverbaende: som over alle Gemeinden onder hun Kreise (spec 14). IT-
    Zweckverbaende: alleen gemeentelijke leden (spec 3.4); een Kreis als lid voegt niets toe."""
    from tam import hand_tables_de as ht
    rows, members = [], []
    for name, r in ht.DE_BEZIRKE.items():
        pop = r.get('population')
        rows.append(new_entity(
            entity_id=r['id'], name=name, match_name=r.get('match_name', name), country='DE',
            kind='de_bezirk', region=r['stadt'], population=pop,
            population_method='own' if pop is not None else 'unknown',
            population_source=r['bron'] if pop is not None else '',
            population_date=str(r.get('peildatum', '')), source=r['bron_url'],
            source_date=str(r.get('peildatum', ''))))
    for kind, table, coop in (('de_hoeherer_kommunalverband', ht.DE_HOEHERE_KOMMUNALVERBAENDE, False),
                              ('de_it_zweckverband', ht.DE_IT_ZWECKVERBAENDE, True)):
        for name, r in table.items():
            keys = list(r.get('members') or [])
            sum_keys = keys
            if coop:
                sum_keys, skipped = _municipal_keys(keys, municipal_population)
                if skipped:
                    log.info('Handtabel %s: niet-gemeentelijke leden tellen niet mee: %s',
                             name, ', '.join(skipped))
            pop = None
            if sum_keys:
                pop, n, unresolved = _key_population(sum_keys, municipal_population, verband_members)
                if unresolved:
                    log.warning('Handtabel %s: leden zonder gemeenten in municipal_population: %s',
                                name, ', '.join(unresolved))
                if n == 0:
                    pop = None
            rows.append(new_entity(
                entity_id=r['id'], name=name, match_name=r.get('match_name', name), country='DE',
                kind=kind, region=r.get('region', ''), population=pop,
                population_method='sum of members' if pop is not None else 'unknown',
                population_source='Destatis Gemeindeverzeichnis (som leden)' if pop is not None else '',
                population_date=population_date if pop is not None else '',
                source=r['bron_url'], source_date=str(r.get('peildatum', ''))))
            if coop:
                for k in keys:
                    members.append({'entity_id': r['id'], 'member_entity_id': _member_id(k),
                                    'member_kind': _member_kind(k), 'membership_type': 'cooperation',
                                    'join_date': ''})
    for name, r in ht.DE_BUNDESMINISTERIEN.items():
        rows.append(new_entity(
            entity_id=r['id'], name=name, match_name=r.get('match_name', name), country='DE',
            kind='de_bundesministerium', population_method='country total',
            source=r['bron_url'], source_date=str(r.get('peildatum', ''))))
    bundle = SourceBundle.empty()
    bundle.entities = pd.DataFrame(rows, columns=bundle.entities.columns)
    bundle.memberships = pd.DataFrame(members, columns=bundle.memberships.columns)
    kinds = bundle.entities['kind'].value_counts().to_dict()
    msg = ', '.join(f'{k} {v}' for k, v in sorted(kinds.items()))
    bundle.status.append(SourceStatus('de_hand', True, len(rows), date.today().isoformat(), msg))
    return bundle
