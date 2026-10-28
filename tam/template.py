"""Invoersjabloon (tam_template.xlsx): aanmaken, bijwerken, lezen en valideren."""
import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from tam.config import COUNTRIES, DETAILED_SEGMENTS, SEGMENTS, all_detailed

DEF_COLUMNS = ['company', 'sales_unit', 'core_product', 'country', 'segments',
               'excluded_kinds', 'min_population', 'notes']
OVERRIDE_COLUMNS = ['accountid', 'entity_id', 'reason']
ALWAYS_IGNORED = ['_Other', 'Undefined']
ENTITY_ID_RE = re.compile(r'^(' + '|'.join(COUNTRIES) + r')-[A-Z]+-.+$')


class TemplateError(ValueError):
    """Ongeldige waarden in het sjabloon; bericht bevat alle problemen."""


@dataclass
class TamDefinition:
    company: str
    sales_unit: str
    core_product: str
    country: str
    segments: list[str]
    excluded_kinds: list[str]
    min_population: int | None
    notes: str


def load_products(path: Path, config: dict) -> tuple[pd.DataFrame, list[dict]]:
    raw = pd.read_excel(path, sheet_name=0)[['Company', 'Sales Unit', 'Core Product Family']]
    raw.columns = ['company', 'sales_unit', 'core_product']
    raw = raw.dropna()
    ignore = set(ALWAYS_IGNORED) | set(config.get('ignore_core_products', []))
    raw = raw[raw['company'].isin(config['companies']) & ~raw['core_product'].isin(ignore)]
    counts = (raw.groupby(['company', 'core_product', 'sales_unit']).size()
              .reset_index(name='rows'))
    counts = counts.sort_values(['company', 'core_product', 'rows', 'sales_unit'],
                                ascending=[True, True, False, True])
    main = counts.drop_duplicates(['company', 'core_product'])
    strays = [{'company': r.company, 'core_product': r.core_product,
               'sales_unit': r.sales_unit, 'rows': int(r.rows)}
              for r in counts.drop(main.index).itertuples()]
    df = (main[['company', 'sales_unit', 'core_product']]
          .sort_values(['company', 'sales_unit', 'core_product']).reset_index(drop=True))
    return df, strays


# Naslag voor het blad 'lists': per detailed segment de landen en wat eronder valt.
DETAILED_HELP = {
    'Municipality': ('NL, BE, DE', 'Gemeente / commune / Gemeinde (in DE: verbandsfrei, '
                     'erfuellende Gemeinde of BW-lid van een GVV)'),
    'County-level city': ('DE', 'Kreisfreie Stadt / Stadtkreis (gemeente met Kreis-taken)'),
    'City district': ('NL, BE, DE', 'Stadsdeel Amsterdam, district Antwerpen, Bezirk Berlijn/Hamburg'),
    'Island authority': ('BQ', 'Openbaar lichaam Bonaire, Sint Eustatius, Saba'),
    'Municipal autonomous agency': ('BE', 'Autonoom gemeentebedrijf (AGB)'),
    'District administration': ('SR', 'Districtscommissariaat (10 districten Suriname)'),
    'State': ('DE', 'Bundesland, incl. stadstaten Berlijn, Hamburg, Bremen'),
    'Region / Community': ('BE', 'Gewest of Gemeenschap'),
    'Province': ('NL, BE', 'Provincie'),
    'Administrative district': ('DE', 'Regierungsbezirk (Bezirksregierung)'),
    'County': ('DE', 'Landkreis / Kreis'),
    'Higher municipal association': ('DE', 'Bezirk (Beieren), Landschaftsverband (NRW), LWV Hessen e.d.'),
    'Water authority': ('NL', 'Waterschap / hoogheemraadschap'),
    'Ministry': ('NL, DE, AW, CW, SX, SR', 'Ministerie / Bundesministerium'),
    'National agency': ('NL, BQ, SR', 'Rijksorganisatie (agentschap, ZBO, inspectie), Rijksdienst '
                        'Caribisch Nederland, regering Suriname'),
    'Federal public service': ('BE', 'FOD / POD'),
    'Country government (Caribbean)': ('AW, CW, SX', 'Regering van Aruba, Curacao, Sint Maarten'),
    'Safety region': ('NL', 'Veiligheidsregio'),
    'Police zone': ('BE', 'Politiezone / zone de police'),
    'Emergency services zone': ('BE', 'Hulpverleningszone / zone de secours'),
    'Municipal federation': ('DE', 'Amt, Samtgemeinde, Verbandsgemeinde, Verwaltungsgemeinschaft, '
                             'Gemeindeverwaltungsverband'),
    'Shared service organization': ('NL, BE', 'Ambtelijke werkorganisatie / shared service center '
                                    '(bv. Samenwerking Kempengemeenten)'),
    'IT cooperation': ('NL, BE, DE', 'ICT-samenwerking, kommunaler IT-Zweckverband'),
    'Tax cooperation': ('NL, BE', 'Belastingsamenwerking (heffing en invordering, WOZ)'),
    'Environmental agency': ('NL, BE', 'Omgevingsdienst'),
    'Public health service (GGD)': ('NL', 'GGD'),
    'Social employment & work': ('NL, BE', 'Werkbedrijf, sociale werkvoorziening, participatiebedrijf'),
    'Waste management': ('NL, BE', 'Afvalsamenwerking / afvalintercommunale'),
    'Recreation authority': ('NL, BE', 'Recreatieschap'),
    'Procurement cooperation': ('NL, BE', 'Inkoopsamenwerking'),
    'Transport authority': ('NL, BE', 'Vervoerregio, OV- of mobiliteitssamenwerking'),
    'Regional cooperation': ('NL, BE', 'Regionale samenwerking van gemeenten (bv. regio, '
                             'streekverband)'),
    'Welfare association': ('BE', 'Welzijnsvereniging (Vlaanderen)'),
    'Special-purpose association': ('DE', 'Zweckverband (water, afval, scholen e.d.)'),
    'Other joint arrangement': ('NL, BE', 'Overige samenwerking (archief, bedrijventerrein, '
                                'haven, jeugd e.d.)'),
}
# Een lichte vulkleur per segment, zodat de groepen in het blad meteen opvallen.
_SEGMENT_FILLS = ['DDEBF7', 'E2EFDA', 'FFF2CC', 'FCE4D6', 'EDE7F6']


def _lists_sheet(wb: Workbook) -> None:
    """(Her)bouw het hulpblad 'lists': een rij per detailed segment met het segment
    ernaast, de landen en wat eronder valt. Bij --update vervangen zodat nieuwe waarden
    erin staan. Het blad is alleen naslag, niet ingevuld."""
    from openpyxl.styles import Font, PatternFill

    if 'lists' in wb.sheetnames:
        del wb['lists']
    ws = wb.create_sheet('lists')
    ws.append(['segment', 'detailed_segment', 'landen', 'wat valt eronder'])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, segment in enumerate(SEGMENTS):
        fill = PatternFill('solid', fgColor=_SEGMENT_FILLS[i % len(_SEGMENT_FILLS)])
        for detailed in DETAILED_SEGMENTS[segment]:
            landen, uitleg = DETAILED_HELP.get(detailed, ('', ''))
            ws.append([segment, detailed, landen, uitleg])
            for cell in ws[ws.max_row]:
                cell.fill = fill
    last = ws.max_row
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:D{last}'
    for col, width in zip('ABCD', (28, 32, 22, 80)):
        ws.column_dimensions[col].width = width
    ws.append([])
    ws.append(['segments: alleen de segmentnamen uit kolom A. excluded_kinds: waarden uit kolom B '
               'die binnen die segmenten vallen.'])
    ws.append(['Meerdere waarden in een cel: scheiden met ; (puntkomma).'])


def _product_rows(products: pd.DataFrame):
    for r in products.itertuples():
        for country in COUNTRIES:
            yield [r.company, r.sales_unit, r.core_product, country, None, None, None, None]


def make_template(products: pd.DataFrame, path: Path, *, update: bool = False) -> int:
    path = Path(path)
    if path.exists() and not update:
        raise FileExistsError(f'{path} bestaat al; gebruik update om rijen toe te voegen')
    if not path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = 'tam_definitions'
        ws.append(DEF_COLUMNS)
        rows = list(_product_rows(products))
        for row in rows:
            ws.append(row)
        wb.create_sheet('match_overrides').append(OVERRIDE_COLUMNS)
        _lists_sheet(wb)
        wb.save(path)
        return len(rows)
    wb = load_workbook(path)
    ws = wb['tam_definitions']
    existing = {(r[0], r[2], r[3]) for r in ws.iter_rows(min_row=2, values_only=True)}
    added = 0
    for row in _product_rows(products):
        if (row[0], row[2], row[3]) not in existing:
            ws.append(row)
            added += 1
    _lists_sheet(wb)
    wb.save(path)
    return added


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == ''


def _split(v) -> list[str]:
    return [] if _blank(v) else [p.strip() for p in str(v).split(';') if p.strip()]


def read_template(path: Path) -> tuple[list[TamDefinition], pd.DataFrame]:
    errors: list[str] = []
    defs: list[TamDefinition] = []
    df = pd.read_excel(path, sheet_name='tam_definitions', dtype=object)
    seen: dict[tuple, int] = {}   # (company, core_product, country) -> eerste Excel-rij
    for idx, row in df.iterrows():
        xl = idx + 2
        segments = _split(row.get('segments'))
        if not segments:
            continue
        key = tuple('' if _blank(row.get(c)) else str(row.get(c)).strip()
                    for c in ('company', 'core_product', 'country'))
        if key in seen:
            # M4: dubbele definitie zou TAM-rijen dubbel tellen
            errors.append(f"tam_definitions rij {xl}: dubbele definitie voor {' / '.join(key)} "
                          f"(ook in rij {seen[key]})")
        else:
            seen[key] = xl
        excluded = _split(row.get('excluded_kinds'))

        def bad(value, column):
            errors.append(f"tam_definitions rij {xl}: onbekende waarde '{value}' in {column}")

        for s in segments:
            if s not in SEGMENTS:
                bad(s, 'segments')
        for k in excluded:
            if k not in all_detailed():
                bad(k, 'excluded_kinds')
        country = str(row.get('country')).strip()
        if country not in COUNTRIES:
            bad(country, 'country')
        minpop = None
        mp = row.get('min_population')
        if not _blank(mp):
            try:
                f = float(mp)
                if f != int(f) or f < 0:
                    raise ValueError
                minpop = int(f)
            except (ValueError, TypeError, OverflowError):
                bad(mp, 'min_population')
        notes = '' if _blank(row.get('notes')) else str(row['notes'])
        defs.append(TamDefinition(str(row['company']), str(row['sales_unit']),
                                  str(row['core_product']), country, segments, excluded,
                                  minpop, notes))
    overrides = pd.read_excel(path, sheet_name='match_overrides', dtype=str)
    overrides = overrides.reindex(columns=OVERRIDE_COLUMNS)
    for idx, eid in overrides['entity_id'].items():
        if not _blank(eid) and not ENTITY_ID_RE.match(str(eid).strip()):
            errors.append(f"match_overrides rij {idx + 2}: onbekende waarde '{eid}' in entity_id")
    if errors:
        raise TemplateError('\n'.join(errors))
    overrides['entity_id'] = overrides['entity_id'].fillna('').str.strip()
    return defs, overrides
