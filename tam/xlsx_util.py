"""Kleine Excel-hulpjes voor template en uitvoer."""
import math

from openpyxl.utils import get_column_letter

MAX_WIDTH = 80   # bovengrens zodat een enkele lange kolom het blad niet onleesbaar maakt


def autofit_columns(ws, quantile: float = 0.8, skip: set | None = None) -> None:
    """Zet elke kolombreedte zo dat minstens `quantile` van de gevulde waarden past,
    en nooit smaller dan de kop. Kolommen in `skip` (letters) houden hun breedte."""
    skip = skip or set()
    for col_idx, cells in enumerate(ws.iter_cols(values_only=True), start=1):
        letter = get_column_letter(col_idx)
        if letter in skip or not cells:
            continue
        header = len(str(cells[0])) if cells[0] is not None else 0
        lengths = sorted(len(str(v)) for v in cells[1:] if v is not None and str(v) != '')
        fit = 0
        if lengths:
            fit = lengths[max(0, math.ceil(quantile * len(lengths)) - 1)]
        ws.column_dimensions[letter].width = min(max(header, fit, 6) + 2, MAX_WIDTH)


def write_instructions(wb, rows: list, title: str = 'instructies') -> None:
    """(Her)bouw een instructieblad als eerste tabblad en maak het actief. `rows` is een
    lijst (kop, tekst); koppen zonder inspringing worden vet, de eerste rij groter."""
    from openpyxl.styles import Alignment, Font

    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, 0)
    for left, right in rows:
        ws.append([left, right])
        row = ws.max_row
        ws.cell(row, 1).alignment = Alignment(vertical='top')
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical='top')
        if left and not left.startswith('  '):
            ws.cell(row, 1).font = Font(bold=True, size=14 if row == 1 else 11)
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 100
    wb.active = 0
