"""Caribisch deel van het Koninkrijk en Suriname (spec sectie 15)."""
import math
import re

import pandas as pd
from openpyxl import Workbook, load_workbook

from tam import hand_tables_carib as hc
from tam.config import COUNTRIES, CRM_COUNTRY_TO_TAM, DETAILED_SEGMENTS, KIND_TO_SEGMENT
from tam.entities import ALIAS_COLUMNS, ENTITY_COLUMNS, new_entity
from tam.master import build_master
from tam.matching import TYPE_TO_KINDS, build_name_index, crm_in_scope, pass1
from tam.sources.carib import load_carib_sr
from tam.sources.nl import load_nl_hand
from tam.template import DEF_COLUMNS, make_template, read_template

NEW = ['AW', 'CW', 'SX', 'BQ', 'SR']
NATIONAL_KINDS = {'carib_land', 'carib_ministerie', 'bq_rijksdienst', 'sr_regering', 'sr_ministerie'}


# ---------- config ----------

def test_new_countries_and_crm_mapping():
    assert COUNTRIES == ['NL', 'BE', 'DE'] + NEW
    for c in ('AW', 'CW', 'SX', 'BQ'):
        assert CRM_COUNTRY_TO_TAM[c] == c


def test_district_administration_appended_to_local_government():
    local = DETAILED_SEGMENTS['Local government']
    assert local[:5] == ['Municipality', 'County-level city', 'City district', 'Island authority',
                         'Municipal autonomous agency']
    assert local[5:] == ['District administration']


def test_new_kinds_mapped():
    assert KIND_TO_SEGMENT['carib_land'] == ('National government', 'Country government (Caribbean)')
    assert KIND_TO_SEGMENT['carib_ministerie'] == ('National government', 'Ministry')
    assert KIND_TO_SEGMENT['bq_openbaar_lichaam'] == ('Local government', 'Island authority')
    assert KIND_TO_SEGMENT['bq_rijksdienst'] == ('National government', 'National agency')
    assert KIND_TO_SEGMENT['sr_regering'] == ('National government', 'National agency')
    assert KIND_TO_SEGMENT['sr_ministerie'] == ('National government', 'Ministry')
    assert KIND_TO_SEGMENT['sr_district'] == ('Local government', 'District administration')
    assert 'nl_land_caribisch' not in KIND_TO_SEGMENT
    assert 'nl_openbaar_lichaam' not in KIND_TO_SEGMENT


# ---------- handtabellen ----------

def _rows():
    for table in hc.TABLES:
        for name, row in table.items():
            yield name, row


def test_every_hand_row_has_source_date_and_fixed_id():
    ids = []
    for name, row in _rows():
        assert str(row.get('bron_url', '')).startswith('https://'), name
        assert row.get('peildatum'), name
        assert 'population' in row, name
        assert row['id'] == row['id'].upper(), name
        assert row['id'].split('-')[0] == row['country'], name
        assert row['id'].startswith(f"{row['country']}-HAND-"), name
        assert row['kind'] in KIND_TO_SEGMENT, name
        ids.append(row['id'])
    assert len(ids) == len(set(ids))


def test_population_has_iso_reference_date():
    for name, row in _rows():
        assert re.fullmatch(r'\d{4}-\d{2}-\d{2}', row['peildatum']), name
        if row['population'] is not None:
            assert isinstance(row['population'], int) and row['population'] > 0, name


def test_suriname_districts():
    districts = {n: r for n, r in _rows() if r['kind'] == 'sr_district'}
    assert len(districts) == 10
    assert sum(r['population'] for r in districts.values()) == 541638


def test_bq_islands_sum():
    islands = {n: r for n, r in _rows() if r['kind'] == 'bq_openbaar_lichaam'}
    assert len(islands) == 3
    assert sum(r['population'] for r in islands.values()) == 33229


# ---------- lader ----------

def test_load_carib_sr_counts_and_methods():
    b = load_carib_sr()
    e = b.entities
    assert list(e.columns) == ENTITY_COLUMNS
    assert e['entity_id'].is_unique
    assert set(e['country']) == set(NEW)
    counts = e.groupby(['country', 'kind']).size().to_dict()
    assert counts[('AW', 'carib_land')] == 1 and counts[('CW', 'carib_land')] == 1
    assert counts[('SX', 'carib_land')] == 1
    assert counts[('BQ', 'bq_openbaar_lichaam')] == 3 and counts[('BQ', 'bq_rijksdienst')] == 1
    assert counts[('SR', 'sr_regering')] == 1 and counts[('SR', 'sr_district')] == 10
    assert counts[('SR', 'sr_ministerie')] == 17
    by_id = e.set_index('entity_id')
    assert by_id.loc['CW-HAND-CURACAO', 'population'] == 158006
    assert by_id.loc['CW-HAND-CURACAO', 'population_method'] == 'own'
    assert by_id.loc['BQ-HAND-BONAIRE', 'official_ids'] == 'cbs:GM9001'
    assert by_id.loc['BQ-HAND-RCN', 'population_method'] == 'country total'
    ministries = e[e['kind'].isin(['carib_ministerie', 'sr_ministerie'])]
    assert set(ministries['population_method']) == {'country total'}
    assert set(e.loc[e['kind'] == 'sr_district', 'detailed_segment']) == {'District administration'}
    assert list(b.aliases.columns) == ALIAS_COLUMNS
    assert set(b.aliases['entity_id']) <= set(e['entity_id'])
    assert (b.aliases.merge(e[['entity_id', 'country']], on='entity_id', suffixes=('', '_e'))
            .eval('country == country_e').all())
    assert b.status[0].name == 'carib_sr' and b.status[0].ok


def test_missing_population_is_unknown_and_logged(caplog):
    with caplog.at_level('WARNING'):
        b = load_carib_sr()
    unknown = b.entities[(b.entities['population_method'] == 'unknown')]
    for eid in unknown['entity_id']:
        assert eid in caplog.text
    assert unknown['population'].isna().all()


def test_nl_hand_has_no_caribbean_rows():
    e = load_nl_hand().entities
    assert set(e['kind']) == {'nl_stadsdeel'}
    assert not e['entity_id'].str.contains('BONAIRE|SABA|EUSTATIUS|ARUBA|CURACAO|MAARTEN').any()


# ---------- landtotalen ----------

def test_country_totals_in_master():
    result = build_master(None, loaders=[('carib_sr', lambda ctx: load_carib_sr())])
    e = result.entities.set_index('entity_id')
    sr = e[e['country'] == 'SR']
    districts = sr[sr['kind'] == 'sr_district']['population'].sum()
    assert districts == 541638
    assert (sr.loc[sr['kind'].isin(['sr_regering', 'sr_ministerie']), 'population'] == districts).all()
    assert e.loc['BQ-HAND-RCN', 'population'] == 33229
    cw = e[(e['country'] == 'CW') & (e['kind'] == 'carib_ministerie')]
    assert len(cw) and (cw['population'] == 158006).all()
    sx = e[(e['country'] == 'SX') & (e['kind'] == 'carib_ministerie')]
    assert len(sx) and (sx['population'] == e.loc['SX-HAND-SINT-MAARTEN', 'population']).all()


def test_country_total_missing_logged_once(caplog):
    from tam.entities import apply_country_totals
    rows = [dict(entity_id=f'AW-HAND-MIN-{i}', name=str(i), country='AW', kind='carib_ministerie',
                 population_method='country total') for i in range(3)]
    ents = pd.DataFrame([new_entity(**r) for r in rows], columns=ENTITY_COLUMNS)
    warnings = []
    with caplog.at_level('WARNING'):
        out = apply_country_totals(ents, {}, warnings)
    records = [r for r in caplog.records if r.name == 'tam.entities']
    assert len(records) == 1 and 'blijft leeg' not in records[0].getMessage()
    assert all(f'AW-HAND-MIN-{i}' in records[0].getMessage() for i in range(3))
    assert warnings == [records[0].getMessage()]
    assert set(out['population_method']) == {'unknown'}


def test_country_total_missing_flagged_and_unknown():
    # land zonder bevolkingscijfer: ministeries leeg, methode unknown, regel in de runlog
    rows = [dict(entity_id='AW-HAND-X', name='X', country='AW', kind='carib_land', population=None,
                 population_method='unknown'),
            dict(entity_id='AW-HAND-MIN-Y', name='Y', country='AW', kind='carib_ministerie',
                 population_method='country total')]
    from tam.entities import SourceBundle, SourceStatus
    b = SourceBundle.empty()
    b.entities = pd.DataFrame([new_entity(**r) for r in rows], columns=ENTITY_COLUMNS)
    b.status.append(SourceStatus('x', True, 2, '', ''))
    result = build_master(None, loaders=[('x', lambda ctx: b)])
    e = result.entities.set_index('entity_id')
    assert math.isnan(e.loc['AW-HAND-MIN-Y', 'population'])
    assert e.loc['AW-HAND-MIN-Y', 'population_method'] == 'unknown'
    assert sum('AW-HAND-MIN-Y' in w for w in result.extras['warnings']) == 1


# ---------- matching ----------

CRM_COLUMNS = ['accountid', 'name', 'canonical_name', 'detected_type', 'detected_country',
               'Segment (category)', 'segment_confidence', 'cx_kvknummer', 'cx_address1_country',
               'cx_address3_country']


def _crm(*rows):
    out = []
    for i, r in enumerate(rows):
        base = {c: float('nan') for c in CRM_COLUMNS}
        base.update({'accountid': f'acc-{i}', 'detected_type': 'onbekend',
                     'Segment (category)': 'Governmental', 'segment_confidence': 'high'})
        base.update(r)
        out.append(base)
    return pd.DataFrame(out, columns=CRM_COLUMNS)


def test_scope_maps_caribbean_to_itself_and_suriname_by_address():
    frame = _crm(
        dict(name='A', detected_country='AW'), dict(name='C', detected_country='CW'),
        dict(name='S', detected_country='SX'), dict(name='B', detected_country='BQ'),
        dict(name='Ministerie van Defensie', detected_country='OTHER', detected_type='ministerie',
             cx_address1_country='Suriname'),
        dict(name='District X', detected_country='OTHER', cx_address3_country=' suriname '),
        dict(name='Ville de Paris', detected_country='OTHER', cx_address1_country='Frankrijk'),
    )
    scope, out_gov = crm_in_scope(frame)
    assert dict(zip(scope['name'], scope['tam_country'])) == {
        'A': 'AW', 'C': 'CW', 'S': 'SX', 'B': 'BQ', 'Ministerie van Defensie': 'SR', 'District X': 'SR'}
    assert list(out_gov['crm_name']) == ['Ville de Paris']


def test_scope_without_address_columns():
    frame = _crm(dict(name='A', detected_country='OTHER')).drop(
        columns=['cx_address1_country', 'cx_address3_country'])
    scope, out_gov = crm_in_scope(frame)
    assert len(scope) == 0 and len(out_gov) == 1


def test_type_to_kinds_include_new_national_kinds():
    for t in ('land', 'ministerie', 'rijksoverheid'):
        assert NATIONAL_KINDS <= TYPE_TO_KINDS[t], t
    assert 'bq_openbaar_lichaam' in TYPE_TO_KINDS['land']
    assert 'nl_ministerie' in TYPE_TO_KINDS['ministerie']


def test_caribbean_ministry_rule_match_in_own_country():
    b = load_carib_sr()
    entities, aliases = b.entities, b.aliases
    frame = _crm(
        dict(name='Ministerie van Justitie', canonical_name='van Justitie', detected_type='ministerie',
             detected_country='CW'),
        dict(name='Land Curaçao', canonical_name='Curaçao', detected_type='land', detected_country='CW'),
        dict(name='Ministerie van Defensie', canonical_name='van Defensie', detected_type='ministerie',
             detected_country='OTHER', cx_address1_country='Suriname'),
        dict(name='Openbaar lichaam Bonaire', canonical_name='Caribisch Nederland', detected_type='land',
             detected_country='BQ'),
    )
    scope, _ = crm_in_scope(frame)
    matched, rest = pass1(scope, entities, build_name_index(entities, aliases))
    assert dict(zip(matched['accountid'], matched['entity_id'])) == {
        'acc-0': 'CW-HAND-MIN-JUSTITIE', 'acc-1': 'CW-HAND-CURACAO',
        'acc-2': 'SR-HAND-MIN-DEFENSIE', 'acc-3': 'BQ-HAND-BONAIRE'}
    assert len(rest) == 0


def test_caribbean_ministry_never_matches_nl_ministry():
    # isolatie: een CW-account 'Ministerie van Justitie' mag niet op een NL-ministerie vallen,
    # ook niet als CW zo'n ministerie niet kent (dan pass 2)
    entities = pd.DataFrame([new_entity(**r) for r in [
        dict(entity_id='NL-OO-1', name='Ministerie van Justitie', country='NL', kind='nl_ministerie'),
        dict(entity_id='NL-OO-2', name='Ministerie van Financiën', country='NL', kind='nl_ministerie'),
        dict(entity_id='CW-HAND-MIN-JUSTITIE', name='Ministerie van Justitie', country='CW',
             kind='carib_ministerie'),
    ]], columns=ENTITY_COLUMNS)
    frame = _crm(
        dict(name='Ministerie van Justitie', canonical_name='van Justitie', detected_type='ministerie',
             detected_country='CW'),
        dict(name='Ministerie van Financiën', canonical_name='van Financiën', detected_type='ministerie',
             detected_country='CW'),
    )
    scope, _ = crm_in_scope(frame)
    alias_frame = pd.DataFrame(columns=ALIAS_COLUMNS)
    matched, rest = pass1(scope, entities, build_name_index(entities, alias_frame))
    assert dict(zip(matched['accountid'], matched['entity_id'])) == {'acc-0': 'CW-HAND-MIN-JUSTITIE'}
    assert list(rest['accountid']) == ['acc-1']


# ---------- sjabloon ----------

def _old_template(path):
    """Sjabloon van voor sectie 15: alleen NL/BE/DE, NL-rij ingevuld, een override."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'tam_definitions'
    ws.append(DEF_COLUMNS)
    for c in ('NL', 'BE', 'DE'):
        ws.append(['Company A', 'Unit 1', 'Product X', c,
                   'Local government' if c == 'NL' else None, None, None, 'hi' if c == 'NL' else None])
    wb.create_sheet('match_overrides').append(['accountid', 'entity_id', 'reason'])
    wb['match_overrides'].append(['123', 'NL-GM-0363', 'manual'])
    ls = wb.create_sheet('lists')
    ls.append(['segments', 'detailed_segments'])
    ls.append(['Local government', 'Municipality'])
    wb.save(path)


def test_update_adds_only_new_country_rows(tmp_path):
    path = tmp_path / 'tpl.xlsx'
    _old_template(path)
    products = pd.DataFrame([('Company A', 'Unit 1', 'Product X')],
                            columns=['company', 'sales_unit', 'core_product'])
    assert make_template(products, path, update=True) == 5
    df = pd.read_excel(path, sheet_name='tam_definitions')
    assert list(df['country']) == ['NL', 'BE', 'DE'] + NEW
    defs, overrides = read_template(path)
    assert len(defs) == 1 and defs[0].country == 'NL' and defs[0].notes == 'hi'
    assert list(overrides['accountid']) == ['123']
    # lijstenblad ververst met de nieuwe detailed-waarde
    wb = load_workbook(path)
    detailed = [r[1] for r in wb['lists'].iter_rows(min_row=2, values_only=True)]
    assert 'District administration' in detailed
    assert make_template(products, path, update=True) == 0


def test_template_accepts_new_countries_and_override_ids(tmp_path):
    path = tmp_path / 'tpl.xlsx'
    products = pd.DataFrame([('Company A', 'Unit 1', 'Product X')],
                            columns=['company', 'sales_unit', 'core_product'])
    make_template(products, path)
    wb = load_workbook(path)
    ws = wb['tam_definitions']
    for row in ws.iter_rows(min_row=2):
        row[4].value = 'National government; Local government'
        row[5].value = 'District administration'
    wb['match_overrides'].append(['1', 'SR-HAND-DISTRICT-WANICA', 'x'])
    wb['match_overrides'].append(['2', 'AW-HAND-ARUBA', 'x'])
    wb.save(path)
    defs, overrides = read_template(path)
    assert [d.country for d in defs] == COUNTRIES
    assert list(overrides['entity_id']) == ['SR-HAND-DISTRICT-WANICA', 'AW-HAND-ARUBA']
