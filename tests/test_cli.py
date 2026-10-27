"""Tests voor de CLI tam_build.py (geen netwerk, geen API)."""
import json
import os

import pandas as pd
import pytest
from openpyxl import load_workbook

import tam_build
from tam.entities import (ALIAS_COLUMNS, ENTITY_COLUMNS, MEMBERSHIP_COLUMNS, SourceBundle,
                          SourceStatus, new_entity)

CONFIG = {'companies': ['Company A'], 'ignore_core_products': []}


def _products(tmp_path):
    path = tmp_path / 'products.xlsx'
    pd.DataFrame([('Company A', 'Unit 1', 'Product X'), ('Company A', 'Unit 1', 'Product Y'),
                  ('Company Z', 'Unit 9', 'Product Q')],
                 columns=['Company', 'Sales Unit', 'Core Product Family']).to_excel(path, index=False)
    return path


def _config(tmp_path):
    path = tmp_path / 'cfg.json'
    path.write_text(json.dumps(CONFIG), encoding='utf-8')
    return path


def _fake_master():
    ents = pd.DataFrame([
        new_entity(entity_id='NL-CBS-GM0001', name='Aadorp', country='NL', kind='nl_gemeente',
                   population=12000, population_method='own'),
        new_entity(entity_id='NL-CBS-GM0002', name='Bedorp', country='NL', kind='nl_gemeente',
                   population=None, population_method='unknown'),
        new_entity(entity_id='NL-OO-1', name='Werkverband Zuid', country='NL',
                   kind='nl_samenwerking', population_method='sum of members'),
    ], columns=ENTITY_COLUMNS)
    mem = pd.DataFrame([['NL-OO-1', 'NL-CBS-GM0001', 'nl_gemeente', 'cooperation', ''],
                        ['NL-OO-1', 'NL-CBS-GM0002', 'nl_gemeente', 'cooperation', '']],
                       columns=MEMBERSHIP_COLUMNS)
    status = [SourceStatus('nl_cbs', True, 2, '2026-01-01', ''),
              SourceStatus('de_gv', False, 0, '', 'download mislukt')]
    return SourceBundle(ents, mem, pd.DataFrame(columns=ALIAS_COLUMNS), status,
                        {'warnings': ['master-waarschuwing']})


def _acc(accountid, name, detected_type='gemeente_nl', country='NL'):
    return {'accountid': accountid, 'name': name, 'detected_type': detected_type,
            'detected_country': country, 'Segment (category)': 'Governmental',
            'segment_confidence': 'high'}


def _crm_csv(tmp_path, extra=()):
    path = tmp_path / 'crm.csv'
    pd.DataFrame([
        {'accountid': 'acc-1', 'name': 'Gemeente Aadorp', 'detected_type': 'gemeente_nl',
         'detected_country': 'NL', 'Segment (category)': 'Governmental', 'segment_confidence': 'high'},
        {'accountid': 'acc-2', 'name': 'Gemeente Bedorp', 'detected_type': 'gemeente_nl',
         'detected_country': 'NL', 'Segment (category)': 'Governmental', 'segment_confidence': 'high'},
        {'accountid': 'acc-3', 'name': 'Stad Elders', 'detected_type': 'onbekend',
         'detected_country': 'OTHER', 'Segment (category)': 'Governmental', 'segment_confidence': 'high'},
        *extra,
    ]).to_csv(path, index=False)
    return path


def _template(tmp_path, *, fill=True, bad=False, overrides=()):
    path = tmp_path / 'tpl.xlsx'
    assert tam_build.main(['--make-template', '--products', str(_products(tmp_path)),
                           '--config', str(_config(tmp_path)), '--template', str(path)]) == 0
    wb = load_workbook(path)
    ws = wb['tam_definitions']
    if fill:
        # rij 2 = Company A / Product X / NL
        ws.cell(2, 5).value = 'Bogus segment' if bad else 'Local government'
        ws.cell(2, 7).value = 10000
    for row in overrides:
        wb['match_overrides'].append(list(row))
    wb.save(path)
    return path


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv('ANTHROPIC_API_KEY', raising=False)
    monkeypatch.setattr(tam_build, 'ENV_PATH', tmp_path / 'geen.env')
    monkeypatch.setattr(tam_build, 'setup_file_logging', lambda name: None)
    calls = {'master': 0}

    def fake_master(cache):
        calls['master'] += 1
        return _fake_master()

    def no_client(*a, **k):
        raise AssertionError('geen Anthropic-client verwacht')

    monkeypatch.setattr(tam_build, 'build_master', fake_master)
    monkeypatch.setattr(tam_build.anthropic, 'Anthropic', no_client)
    return calls


def _run_args(tmp_path, tpl, *extra, crm_extra=()):
    return ['--template', str(tpl), '--crm', str(_crm_csv(tmp_path, crm_extra)),
            '--products', str(tmp_path / 'products.xlsx'), '--config', str(tmp_path / 'cfg.json'),
            '--output', str(tmp_path / 'out.xlsx'), *extra]


def test_make_template_end_to_end(tmp_path):
    tpl = tmp_path / 'tpl.xlsx'
    rc = tam_build.main(['--make-template', '--products', str(_products(tmp_path)),
                         '--config', str(_config(tmp_path)), '--template', str(tpl)])
    assert rc == 0
    df = pd.read_excel(tpl, sheet_name='tam_definitions')
    assert len(df) == 3 * 2
    assert set(df['company']) == {'Company A'}


def test_make_template_twice_needs_update(tmp_path, capsys):
    args = ['--make-template', '--products', str(_products(tmp_path)),
            '--config', str(_config(tmp_path)), '--template', str(tmp_path / 'tpl.xlsx')]
    assert tam_build.main(args) == 0
    assert tam_build.main(args) == 2
    assert '--update' in capsys.readouterr().err
    assert tam_build.main(args + ['--update']) == 0


def test_missing_config_points_to_example(tmp_path, capsys):
    rc = tam_build.main(['--make-template', '--products', str(_products(tmp_path)),
                         '--config', str(tmp_path / 'nope.json'),
                         '--template', str(tmp_path / 'tpl.xlsx')])
    assert rc == 2
    assert 'tam_config.example.json' in capsys.readouterr().err


def test_template_error_exits_before_master(tmp_path, env, capsys):
    tpl = _template(tmp_path, bad=True, overrides=[('acc-1', 'bad id', '')])
    rc = tam_build.main(_run_args(tmp_path, tpl, '--no-ai'))
    assert rc == 2
    assert env['master'] == 0
    err = capsys.readouterr().err
    assert "'Bogus segment'" in err and "'bad id'" in err


def test_empty_template_exits_cleanly(tmp_path, env, capsys):
    tpl = _template(tmp_path, fill=False)
    rc = tam_build.main(_run_args(tmp_path, tpl, '--dry-run'))
    assert rc == 0
    assert env['master'] == 0
    assert 'Geen TAM-definities' in capsys.readouterr().out
    assert not (tmp_path / 'out.xlsx').exists()


def test_missing_api_key_exits_2(tmp_path, env, capsys):
    tpl = _template(tmp_path)
    rc = tam_build.main(_run_args(tmp_path, tpl))
    assert rc == 2
    assert env['master'] == 0
    assert 'ANTHROPIC_API_KEY' in capsys.readouterr().err


def test_dry_run_writes_no_output(tmp_path, env, capsys):
    tpl = _template(tmp_path)
    rc = tam_build.main(_run_args(tmp_path, tpl, '--no-ai', '--dry-run'))
    assert rc == 0
    assert env['master'] == 1
    assert not (tmp_path / 'out.xlsx').exists()
    assert not (tmp_path / 'reference' / 'tam' / 'known_detailed.json').exists()
    out = capsys.readouterr().out
    assert 'AI-schatting' in out
    assert 'Gemeente Aadorp' not in out


def test_dry_run_without_key_makes_no_ai_call(tmp_path, env, capsys):
    tpl = _template(tmp_path)
    rc = tam_build.main(_run_args(tmp_path, tpl, '--dry-run'))
    assert rc == 0
    out = capsys.readouterr().out
    # een samenwerking zonder regel/cache staat open voor classificatie
    assert '1 nog voor AI' in out
    assert 'plafond $50.00' in out


def test_no_ai_run_writes_output(tmp_path, env, capsys):
    tpl = _template(tmp_path, overrides=[('acc-2', '', 'geen klant')])
    rc = tam_build.main(_run_args(tmp_path, tpl, '--no-ai'))
    assert rc == 0
    out_path = tmp_path / 'out.xlsx'
    assert out_path.exists()
    tam = pd.read_excel(out_path, sheet_name='TAM')
    assert sorted(tam['name']) == ['Aadorp', 'Bedorp']
    status = dict(zip(tam['name'], tam['in_tam']))
    assert status == {'Aadorp': 'yes', 'Bedorp': 'unknown - population missing'}
    assert dict(zip(tam['name'], tam['matched_accounts'])) == {'Aadorp': 1, 'Bedorp': 0}
    matches = pd.read_excel(out_path, sheet_name='CRM matches')
    assert list(matches['accountid']) == ['acc-1']
    assert list(matches['match_pass']) == ['rule']
    unmatched = pd.read_excel(out_path, sheet_name='CRM unmatched')
    reasons = dict(zip(unmatched['accountid'], unmatched['reason']))
    assert reasons == {'acc-2': 'override: geen match', 'acc-3': 'land buiten scope'}
    ents = pd.read_excel(out_path, sheet_name='Entities')
    coop = ents[ents['entity_id'] == 'NL-OO-1'].iloc[0]
    assert coop['detailed_segment'] == 'Other joint arrangement'
    log = pd.read_excel(out_path, sheet_name='Run log', dtype=str).fillna('')
    values = ' | '.join(log['onderwerp'] + ' = ' + log['waarde'])
    assert 'master-waarschuwing' in values
    assert 'NL-OO-1' in values  # niet-geclassificeerde ids als lijst
    assert 'FOUT: download mislukt' in values
    assert 'NL / nl_gemeente = 2' in values
    out = capsys.readouterr().out
    assert 'de_gv' in out and str(out_path) in out
    assert 'Gemeente Aadorp' not in out


def test_env_file_loader(tmp_path, monkeypatch):
    monkeypatch.delenv('ANTHROPIC_API_KEY', raising=False)
    env_file = tmp_path / '.env'
    env_file.write_text('# commentaar\nOTHER=x\nexport ANTHROPIC_API_KEY="sk-test"\n', encoding='utf-8')
    assert tam_build.load_env_key(env_file) is True
    assert os.environ['ANTHROPIC_API_KEY'] == 'sk-test'
    assert 'OTHER' not in os.environ
    monkeypatch.setenv('ANTHROPIC_API_KEY', 'sk-al')
    assert tam_build.load_env_key(env_file) is True
    assert os.environ['ANTHROPIC_API_KEY'] == 'sk-al'
    monkeypatch.delenv('ANTHROPIC_API_KEY')
    assert tam_build.load_env_key(tmp_path / 'missing.env') is False


def test_ai_run_shares_tracker_and_reports_spend(tmp_path, env, monkeypatch, capsys):
    from tam import classify as cl
    tpl = _template(tmp_path)
    monkeypatch.setenv('ANTHROPIC_API_KEY', 'sk-test')
    client = object()
    monkeypatch.setattr(tam_build.anthropic, 'Anthropic', lambda: client)
    seen = {}

    def fake_call(c, tracker, **kw):
        seen['client'], seen['tracker'] = c, tracker
        tracker.spent += 0.25
        return cl.ClassBatch(items=[cl.ClassItem(volgnummer=1, detailed_segment='IT cooperation')])

    monkeypatch.setattr(cl, 'call_parsed', fake_call)
    rc = tam_build.main(_run_args(tmp_path, tpl, '--max-cost', '5'))
    assert rc == 0
    assert seen['client'] is client and seen['tracker'].max_cost == 5.0
    ents = pd.read_excel(tmp_path / 'out.xlsx', sheet_name='Entities')
    assert ents.set_index('entity_id').at['NL-OO-1', 'detailed_segment'] == 'IT cooperation'
    assert 'AI-uitgaven: $0.25' in capsys.readouterr().out


# een account met fuzzy-kandidaat (gaat naar pass 2) en een zonder kandidaat
PASS2_ACCOUNTS = (_acc('acc-4', 'Gemeente Aadorpp'), _acc('acc-5', 'Xyzzy Qwerty', 'onbekend'))


@pytest.mark.filterwarnings('error::FutureWarning')
def test_unmatched_concat_without_future_warning(tmp_path, env):
    tpl = _template(tmp_path)
    rc = tam_build.main(_run_args(tmp_path, tpl, '--no-ai', crm_extra=PASS2_ACCOUNTS))
    assert rc == 0
    unmatched = pd.read_excel(tmp_path / 'out.xlsx', sheet_name='CRM unmatched')
    reasons = dict(zip(unmatched['accountid'], unmatched['reason']))
    assert reasons == {'acc-4': 'AI overgeslagen (--no-ai)',
                       'acc-5': 'geen kandidaat boven drempel',
                       'acc-3': 'land buiten scope'}
    score = dict(zip(unmatched['accountid'], unmatched['best_score']))
    assert score['acc-4'] > 0 and pd.isna(score['acc-5']) and pd.isna(score['acc-3'])


def test_override_to_entity_skips_pass2_and_unmatched(tmp_path, env, monkeypatch, capsys):
    tpl = _template(tmp_path, overrides=[('acc-4', 'NL-CBS-GM0002', 'handmatig'),
                                         ('acc-2', '', 'geen klant')])
    seen = {}
    real_classify, real_pass2 = tam_build.classify_cooperations, tam_build.pass2

    def spy_classify(entities, client, tracker, **kw):
        seen.setdefault('classify_trackers', []).append(tracker)
        return real_classify(entities, client, tracker, **kw)

    def spy_pass2(rest, entities, aliases, client, tracker, **kw):
        seen['pass2_accounts'] = list(rest['accountid'])
        seen['pass2_tracker'] = tracker
        return real_pass2(rest, entities, aliases, client, tracker, **kw)

    monkeypatch.setattr(tam_build, 'classify_cooperations', spy_classify)
    monkeypatch.setattr(tam_build, 'pass2', spy_pass2)
    rc = tam_build.main(_run_args(tmp_path, tpl, '--no-ai', crm_extra=PASS2_ACCOUNTS))
    assert rc == 0
    assert seen['pass2_accounts'] == ['acc-5']
    assert all(t is seen['pass2_tracker'] for t in seen['classify_trackers'])
    out_path = tmp_path / 'out.xlsx'
    matches = pd.read_excel(out_path, sheet_name='CRM matches')
    assert dict(zip(matches['accountid'], matches['match_pass'])) == {'acc-1': 'rule',
                                                                      'acc-4': 'override'}
    unmatched = pd.read_excel(out_path, sheet_name='CRM unmatched')
    assert set(unmatched['accountid']) == {'acc-2', 'acc-3', 'acc-5'}
    log = pd.read_excel(out_path, sheet_name='Run log', dtype=str).fillna('')
    pairs = dict(zip(log['onderwerp'], log['waarde']))
    assert pairs['overrides naar entiteit'] == '1'
    assert pairs['overrides geen match'] == '1'
    assert pairs['classificatie pending'] == '1'


def test_output_not_writable_stops_before_master(tmp_path, env, capsys):
    tpl = _template(tmp_path)
    (tmp_path / '.~lock.out.xlsx#').write_text('x')
    rc = tam_build.main(_run_args(tmp_path, tpl, '--no-ai'))
    assert rc == 2
    assert env['master'] == 0
    assert 'out.xlsx' in capsys.readouterr().err


def test_known_detailed_saved_only_after_write(tmp_path, env, monkeypatch):
    tpl = _template(tmp_path)

    def boom(*a, **k):
        raise OSError('schijf vol')

    monkeypatch.setattr(tam_build, 'write_output', boom)
    with pytest.raises(OSError):
        tam_build.main(_run_args(tmp_path, tpl, '--no-ai'))
    assert not (tmp_path / 'reference' / 'tam' / 'known_detailed.json').exists()


def test_classification_cap_reported(tmp_path, env, monkeypatch, capsys):
    from tam import classify as cl
    from tam.ai import CostCapReached
    tpl = _template(tmp_path)
    monkeypatch.setenv('ANTHROPIC_API_KEY', 'sk-test')
    monkeypatch.setattr(tam_build.anthropic, 'Anthropic', lambda: object())

    def capped(*a, **k):
        raise CostCapReached('plafond')

    monkeypatch.setattr(cl, 'call_parsed', capped)
    rc = tam_build.main(_run_args(tmp_path, tpl))
    assert rc == 0
    out = capsys.readouterr().out
    assert 'Kostenplafond bereikt tijdens classificatie' in out
    assert 'Niet geclassificeerd: 1' in out
    log = pd.read_excel(tmp_path / 'out.xlsx', sheet_name='Run log', dtype=str).fillna('')
    text = ' | '.join(log['waarde'])
    assert 'kostenplafond 1' in text and 'NL-OO-1' in text


def test_override_for_unknown_account_in_run_log_and_not_counted(tmp_path, env):
    tpl = _template(tmp_path, overrides=[('acc-99', 'NL-CBS-GM0001', 'oud account'),
                                         ('acc-2', '', 'geen klant')])
    rc = tam_build.main(_run_args(tmp_path, tpl, '--no-ai'))
    assert rc == 0
    log = pd.read_excel(tmp_path / 'out.xlsx', sheet_name='Run log', dtype=str).fillna('')
    pairs = dict(zip(log['onderwerp'], log['waarde']))
    assert pairs['overrides naar entiteit'] == '0'
    assert pairs['overrides geen match'] == '1'
    assert any('acc-99' in v and 'niet in de CRM' in v for v in log['waarde'])
    matches = pd.read_excel(tmp_path / 'out.xlsx', sheet_name='CRM matches')
    assert 'acc-99' not in set(matches['accountid'])
