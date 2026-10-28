import pandas as pd
from openpyxl import load_workbook

from tam.entities import ENTITY_COLUMNS, MEMBERSHIP_COLUMNS, new_entity
from tam.output import build_run_log, check_known_detailed, write_output
from tam.tam_rows import TAM_ROW_COLUMNS

MATCH_COLS = ['entity_id', 'accountid', 'crm_name', 'match_pass', 'match_reason', 'confidence']
UNM_COLS = ['accountid', 'crm_name', 'tam_country', 'best_candidate', 'best_score', 'reason']


def frames():
    ents = pd.DataFrame([
        new_entity(entity_id='NL-1', name='A', country='NL', kind='nl_gemeente', official_ids='0001234'),
        new_entity(entity_id='NL-GR1', name='GR', country='NL', kind='nl_samenwerking',
                   detailed_segment='IT cooperation'),
    ], columns=ENTITY_COLUMNS)
    mem = pd.DataFrame([
        ['NL-GR1', 'NL-1', 'nl_gemeente', 'cooperation', ''],
        ['NL-P1', 'NL-1', 'nl_gemeente', 'territorial', ''],
        ['NL-1', 'NL-2', 'nl_gemeente', 'same_as', ''],
    ], columns=MEMBERSHIP_COLUMNS)
    return dict(tam=pd.DataFrame(columns=TAM_ROW_COLUMNS), matches=pd.DataFrame(columns=MATCH_COLS),
                entities=ents, memberships=mem, unmatched=pd.DataFrame(columns=UNM_COLS),
                run_log=build_run_log(spend=1.5))


def test_sheet_order_and_names(tmp_path):
    p = tmp_path / 'o.xlsx'
    assert write_output(p, **frames()) == [p]
    wb = load_workbook(p)
    assert wb.sheetnames == ['TAM', 'CRM matches', 'Entities', 'Memberships', 'CRM unmatched', 'Run log']
    for ws in wb.worksheets:
        assert ws.freeze_panes == 'A2' and ws.auto_filter.ref, ws.title
    ws = wb['Entities']
    col = [c.value for c in ws[1]].index('official_ids')
    assert ws.cell(2, col + 1).value == '0001234'


def test_only_cooperation_memberships_exported(tmp_path):
    p = tmp_path / 'o.xlsx'
    write_output(p, **frames())
    df = pd.read_excel(p, sheet_name='Memberships', dtype=str)
    assert len(df) == 1
    assert 'membership_type' not in df.columns
    assert df.iloc[0]['detailed_segment'] == 'IT cooperation'
    assert list(df.columns).index('detailed_segment') == 1


def test_new_detailed_value_reported(tmp_path):
    path = tmp_path / 'known.json'
    assert check_known_detailed({'A', 'B'}, path) == {'A', 'B'}
    assert check_known_detailed({'B', 'C'}, path) == {'C'}
    assert check_known_detailed({'A'}, path) == set()
    assert check_known_detailed({'D'}, path, save=False) == {'D'}
    assert check_known_detailed({'D'}, path, save=False) == {'D'}


def test_match_stats_section():
    log = build_run_log(match_stats={'overrides naar entiteit': 2})
    assert ('== Matching ==', '') in log and ('overrides naar entiteit', '2') in log


def test_run_log_sections():
    log = build_run_log(inputs={'x': 1}, new_detailed={'Z'}, warnings=['w'])
    assert ('== Bronnen ==', '') not in log
    assert ('== Invoer ==', '') in log and ('x', '1') in log
    assert any('Z' in v for _, v in log)


def test_strays_and_branches():
    from tam.entities import SourceStatus
    log = build_run_log(
        strays=[dict(company='C', core_product='P', sales_unit='S', rows=3)],
        entity_counts=pd.DataFrame([dict(country='NL', kind='nl_gemeente', count=342)]),
        tam_status_counts=pd.DataFrame([dict(company='C', sales_unit='S', core_product='P',
                                             country='NL', in_tam='yes', count=5)]),
        population_unknown={'NL': 2}, ai_stats={'calls': 4},
        source_status=[SourceStatus('cbs', True, 10, '2026-01-01', ''),
                       SourceStatus('x', False, 0, '', 'kapot')])
    assert ('C / P', 'S: 3 regels') in log
    assert ('NL / nl_gemeente', 342) in log
    assert ('C / P / NL / yes', 5) in log
    assert ('NL', '2') in log and ('calls', '4') in log
    assert ('cbs', 'ok, 10 rijen, 2026-01-01') in log and ('x', 'FOUT: kapot') in log


def test_bad_frames_raise():
    import pytest
    with pytest.raises(ValueError):
        build_run_log(entity_counts=pd.DataFrame({'a': [1]}))
    with pytest.raises(ValueError):
        build_run_log(tam_status_counts=pd.DataFrame({'a': [1]}))


def test_missing_cooperation_warns(tmp_path, caplog):
    f = frames()
    f['memberships'] = pd.DataFrame([['NL-X', 'NL-1', 'nl_gemeente', 'cooperation', '']],
                                    columns=MEMBERSHIP_COLUMNS)
    write_output(tmp_path / 'o.xlsx', **f)
    assert 'NL-X' in caplog.text


def test_open_points_section():
    log = build_run_log(open_points=['OP-dump: licentie niet vermeld'])
    assert ('== Open punten ==', '') in log
    assert ('open punt', 'OP-dump: licentie niet vermeld') in log
