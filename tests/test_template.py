import pandas as pd
import pytest

from tam.config import COUNTRIES
from tam.template import (TamDefinition, TemplateError, load_products,
                          make_template, read_template)

CONFIG = {'companies': ['Company A', 'Company B'],
          'ignore_core_products': ['Non-product family']}


def _products_xlsx(tmp_path, rows):
    path = tmp_path / 'products.xlsx'
    pd.DataFrame(rows, columns=['Company', 'Sales Unit', 'Core Product Family']).to_excel(
        path, index=False)
    return path


def _products(tmp_path):
    rows = [('Company A', 'Unit 1', 'Product X')] * 3 + [
        ('Company A', 'Unit 2', 'Product X'),
        ('Company B', 'Unit 3', 'Product Y'),
    ]
    return load_products(_products_xlsx(tmp_path, rows), CONFIG)


def test_company_filter_and_ignore_list(tmp_path):
    rows = [
        ('Company A', 'Unit 1', 'Product X'),
        ('Company Z', 'Unit 9', 'Product Q'),
        ('Company A', 'Unit 1', '_Other'),
        ('Company A', 'Unit 1', 'Undefined'),
        ('Company B', 'Unit 3', 'Non-product family'),
        ('Company B', 'Unit 3', 'Product Y'),
    ]
    df, strays = load_products(_products_xlsx(tmp_path, rows), CONFIG)
    assert sorted(df['core_product']) == ['Product X', 'Product Y']
    assert strays == []


def test_stray_assigned_to_main_sales_unit(tmp_path):
    df, strays = _products(tmp_path)
    x = df[df['core_product'] == 'Product X']
    assert list(x['sales_unit']) == ['Unit 1']
    assert strays == [{'company': 'Company A', 'core_product': 'Product X',
                       'sales_unit': 'Unit 2', 'rows': 1}]


def test_make_refuses_overwrite(tmp_path):
    df, _ = _products(tmp_path)
    path = tmp_path / 'tpl.xlsx'
    assert make_template(df, path) == 2 * len(COUNTRIES)
    with pytest.raises(FileExistsError):
        make_template(df, path)


def _fill(path, updates):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb['tam_definitions']
    headers = [c.value for c in ws[1]]
    for row, values in updates.items():
        for col, v in values.items():
            ws.cell(row=row, column=headers.index(col) + 1, value=v)
    wb.save(path)


def test_update_keeps_filled_rows_overrides_and_adds_new(tmp_path):
    df, _ = _products(tmp_path)
    path = tmp_path / 'tpl.xlsx'
    make_template(df, path)
    _fill(path, {2: {'segments': 'Local government; Regional government',
                     'min_population': 10000.0, 'notes': 'hi'}})
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wb['match_overrides'].append(['123', 'NL-GM-0363', 'manual'])
    wb.save(path)

    new_rows = [('Company A', 'Unit 1', 'Product X'), ('Company B', 'Unit 3', 'Product Y'),
                ('Company B', 'Unit 4', 'Product W')]
    df2, _ = load_products(_products_xlsx(tmp_path, new_rows), CONFIG)
    assert make_template(df2, path, update=True) == len(COUNTRIES)

    defs, overrides = read_template(path)
    assert len(defs) == 1
    d = defs[0]
    assert d == TamDefinition('Company A', 'Unit 1', 'Product X', 'NL',
                              ['Local government', 'Regional government'], [], 10000, 'hi')
    assert list(overrides['accountid']) == ['123']
    assert len(pd.read_excel(path, sheet_name='tam_definitions')) == 3 * len(COUNTRIES)


def test_unknown_segment_value_stops_with_row_number(tmp_path):
    df, _ = _products(tmp_path)
    path = tmp_path / 'tpl.xlsx'
    make_template(df, path)
    _fill(path, {3: {'segments': 'Local goverment'},
                 4: {'segments': 'Local government', 'excluded_kinds': 'Nope',
                     'min_population': -5}})
    with pytest.raises(TemplateError) as e:
        read_template(path)
    msg = str(e.value)
    assert 'rij 3' in msg and "'Local goverment'" in msg
    assert "rij 4: onbekende waarde 'Nope' in excluded_kinds" in msg
    assert 'min_population' in msg


def test_empty_segments_row_skipped(tmp_path):
    df, _ = _products(tmp_path)
    path = tmp_path / 'tpl.xlsx'
    make_template(df, path)
    _fill(path, {2: {'segments': 'Local government', 'excluded_kinds': 'Island authority'}})
    defs, overrides = read_template(path)
    assert len(defs) == 1 and defs[0].excluded_kinds == ['Island authority']
    assert overrides.empty


def test_bad_override_entity_id(tmp_path):
    df, _ = _products(tmp_path)
    path = tmp_path / 'tpl.xlsx'
    make_template(df, path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wb['match_overrides'].append(['1', 'bogus', 'x'])
    wb['match_overrides'].append(['2', None, 'no match'])
    wb.save(path)
    with pytest.raises(TemplateError) as e:
        read_template(path)
    assert 'match_overrides rij 2' in str(e.value)
    assert 'rij 3' not in str(e.value)


def test_duplicate_definition_rows_rejected(tmp_path):
    # M4: twee ingevulde rijen voor dezelfde (company, core_product, country) = fout
    df, _ = _products(tmp_path)
    path = tmp_path / 'tpl.xlsx'
    make_template(df, path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wb['tam_definitions'].append(['Company A', 'Unit 1', 'Product X', 'NL', 'Regional government',
                                  None, None, None])
    wb['tam_definitions'].append(['Company A', 'Unit 1', 'Product X', 'BE', None, None, None, None])
    wb.save(path)
    _fill(path, {2: {'segments': 'Local government'}, 3: {'segments': 'Local government'}})
    with pytest.raises(TemplateError) as e:
        read_template(path)
    msg = str(e.value)
    dup_row = 2 + 2 * len(COUNTRIES)   # eerste toegevoegde rij na de kop en alle productrijen
    assert f'rij {dup_row}:' in msg and 'rij 2)' in msg and 'Product X' in msg
    # de lege dubbele rij (BE, zonder segments) telt niet
    assert f'rij {dup_row + 1}' not in msg
