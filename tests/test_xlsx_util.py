from openpyxl import Workbook

from tam.xlsx_util import autofit_columns


def test_width_covers_80_percent_of_values_and_header():
    wb = Workbook()
    ws = wb.active
    ws.append(['id', 'naam'])
    for n in [5, 5, 6, 6, 7, 8, 9, 10, 40, 200]:
        ws.append([1, 'x' * n])
    autofit_columns(ws)
    w = ws.column_dimensions['B'].width
    lengths = sorted([5, 5, 6, 6, 7, 8, 9, 10, 40, 200])
    covered = sum(1 for n in lengths if n <= w)
    assert covered >= 8 and w < 200
    assert ws.column_dimensions['A'].width >= len('id')


def test_mostly_empty_column_uses_filled_values_and_header():
    wb = Workbook()
    ws = wb.active
    ws.append(['segments'])
    for v in [None] * 20 + ['Local government; Inter-municipal cooperation']:
        ws.append([v])
    autofit_columns(ws)
    assert ws.column_dimensions['A'].width >= len('Local government; Inter-municipal cooperation')


def test_skip_keeps_existing_width():
    wb = Workbook()
    ws = wb.active
    ws.append(['a'])
    ws.column_dimensions['A'].width = 99
    autofit_columns(ws, skip={'A'})
    assert ws.column_dimensions['A'].width == 99
